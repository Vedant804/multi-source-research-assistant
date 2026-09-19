import asyncio
import io
import json
import logging
import uuid
from pathlib import Path

import numpy as np
from langchain_text_splitters import RecursiveCharacterTextSplitter
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb

from app import db
from app.config import settings
from app.llm import get_embeddings
from app.redis_client import get_redis
from app.services.data_ops import inspect_tables

logger = logging.getLogger(__name__)

TEXT_EXT = {".pdf", ".docx", ".txt", ".md"}
DATA_EXT = {".csv", ".xlsx"}

_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150)


def _docs_key(session_id: str) -> str:
    return f"session:{session_id}:docs"


def _extract_text(ext: str, content: bytes) -> str:
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(content))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return content.decode("utf-8", errors="ignore")


def _public(row: dict) -> dict:
    meta = row["meta"] or {}
    if row["kind"] == "text":
        summary = {"chunks": meta.get("chunks", 0)}
    else:
        summary = {
            "tables": [
                {"name": t["name"], "rows": t["rows"], "columns": [c["name"] for c in t["columns"]]}
                for t in meta.get("tables", [])
            ]
        }
    created = row["created_at"]
    return {
        "id": str(row["id"]),
        "filename": row["filename"],
        "kind": row["kind"],
        "created_at": created.isoformat() if hasattr(created, "isoformat") else str(created),
        "summary": summary,
    }


async def ingest_file(session_id: str, filename: str, content: bytes) -> dict:
    name = Path(filename).name
    ext = Path(name).suffix.lower()
    if ext not in TEXT_EXT | DATA_EXT:
        raise ValueError(f"Unsupported file type '{ext}'. Allowed: {sorted(TEXT_EXT | DATA_EXT)}")

    doc_id = uuid.uuid4()

    if ext in TEXT_EXT:
        text = await asyncio.to_thread(_extract_text, ext, content)
        chunks = [c for c in _splitter.split_text(text) if c.strip()]
        if not chunks:
            raise ValueError("No extractable text found (scanned PDFs need OCR)")

        emb = get_embeddings()
        vectors: list[list[float]] = []
        for i in range(0, len(chunks), 64):
            vectors += await emb.aembed_documents(chunks[i : i + 64])

        async with db.pool.connection() as conn:
            async with conn.transaction():
                await conn.execute(
                    "INSERT INTO documents (id, session_id, filename, kind, meta) VALUES (%s,%s,%s,'text',%s)",
                    (doc_id, session_id, name, Jsonb({"chunks": len(chunks)})),
                )
                async with conn.cursor() as cur:
                    await cur.executemany(
                        "INSERT INTO chunks (document_id, session_id, chunk_index, content, embedding) "
                        "VALUES (%s,%s,%s,%s,%s)",
                        [(doc_id, session_id, i, c, np.array(v)) for i, (c, v) in enumerate(zip(chunks, vectors))],
                    )
    else:
        path = Path(settings.upload_dir) / session_id / f"{doc_id}{ext}"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        try:
            tables = await asyncio.to_thread(inspect_tables, str(path), name, ext)
        except Exception as exc:
            path.unlink(missing_ok=True)
            raise ValueError(f"Could not parse table file: {exc}")
        async with db.pool.connection() as conn:
            await conn.execute(
                "INSERT INTO documents (id, session_id, filename, kind, meta) VALUES (%s,%s,%s,'data',%s)",
                (doc_id, session_id, name, Jsonb({"path": str(path), "tables": tables})),
            )

    await get_redis().delete(_docs_key(session_id))
    docs = await list_documents(session_id)
    return next(d for d in docs if d["id"] == str(doc_id))


async def list_documents(session_id: str) -> list[dict]:
    r = get_redis()
    cached = await r.get(_docs_key(session_id))
    if cached:
        return json.loads(cached)

    async with db.pool.connection() as conn:
        async with conn.cursor(row_factory=dict_row) as cur:
            await cur.execute(
                "SELECT id, filename, kind, meta, created_at FROM documents "
                "WHERE session_id=%s ORDER BY created_at DESC",
                (session_id,),
            )
            rows = await cur.fetchall()

    docs = [_public(row) for row in rows]
    await r.set(_docs_key(session_id), json.dumps(docs), ex=settings.session_cache_ttl)
    return docs


async def get_data_files(session_id: str) -> list[dict]:
    """Internal: data files with their on-disk path and full table metadata."""
    async with db.pool.connection() as conn:
        async with conn.cursor(row_factory=dict_row) as cur:
            await cur.execute(
                "SELECT id, filename, meta FROM documents WHERE session_id=%s AND kind='data'",
                (session_id,),
            )
            rows = await cur.fetchall()
    return [{"id": str(r["id"]), "filename": r["filename"], **r["meta"]} for r in rows]


async def delete_document(session_id: str, doc_id: str) -> bool:
    try:
        uid = uuid.UUID(doc_id)
    except ValueError:
        return False
    async with db.pool.connection() as conn:
        async with conn.cursor(row_factory=dict_row) as cur:
            await cur.execute(
                "DELETE FROM documents WHERE id=%s AND session_id=%s RETURNING kind, meta",
                (uid, session_id),
            )
            row = await cur.fetchone()
    if not row:
        return False
    if row["kind"] == "data":
        Path(row["meta"].get("path", "")).unlink(missing_ok=True)
    await get_redis().delete(_docs_key(session_id))
    return True